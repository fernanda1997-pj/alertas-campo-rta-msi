"""Gera js/modelo-ficha.js e modelos/ficha_<R>_<pav|np>.xlsx a partir das
Fichas de Inspeção reais mais recentes de cada região + tipo.

- js/modelo-ficha.js: as linhas de cada ficha (trecho → S.R.E. → km início/fim),
  na ordem da própria ficha, e ONDE fica cada uma no Excel (aba, linha, coluna
  de cada opção de cada grupo). Usado pelo CSV "Linhas da ficha" do Painel e
  pela aba "Ficha" (preencher em campo e baixar a planilha .xlsx preenchida).
- modelos/*.xlsx: cópia da ficha real usada como molde. Na exportação o app
  apaga os X das linhas de dados e marca os novos direto no XML da planilha —
  formatação, fórmulas e gráficos ficam intactos.

Por que a ficha e não o cadastro (CADASTRO_SRE_TRECHOS_GEORREF_REV01.xlsx) nem
o shapefile: em 2026-10-29 os três não batiam (ex.: 201ETO0250 é do trecho 19
no cadastro e do 20 na ficha; 201ETO0270 tem 12,468 km no cadastro e 12,570 na
ficha). Como o resultado é pra entregar NA FICHA, quem manda é ela.

Usa os helpers do converter das fichas (web - fichas/ficha-inspecao/converter_fichas.py)
pra achar cabeçalho, grupos e colunas — os mesmos que o geoportal usa.
Rodar de novo quando chegar ficha de um mês novo (ou o modelo mudar):

    python ferramentas/gerar_modelo_ficha.py
"""
import contextlib
import glob
import io
import json
import os
import re
import shutil
import sys

import openpyxl
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
FICHAS = os.path.join(RAIZ, '..', 'web - fichas', 'ficha-inspecao')
sys.path.insert(0, FICHAS)

with contextlib.redirect_stdout(io.StringIO()):
    import converter_fichas as cf  # noqa: E402

SRE_OK = re.compile(r'^(\d{3}ETO\d{4}[A-Z]?|ACESSO\d+)$')


def tipo_do_arquivo(nome):
    n = cf._norm(os.path.basename(nome)).replace(' ', '_')
    return 'np' if 'NAO_PAVIMENTADA' in n else 'pav'


def regiao_e_mes(ws_list, nome_arquivo):
    """Região (R1..R13) e competência (AAAA-MM) da ficha, lidas do cabeçalho
    da 1ª aba de trecho ou, se faltar, do nome do arquivo."""
    reg = comp = None
    for ws in ws_list:
        for r in range(1, 5):
            for c in range(1, min(ws.max_column, 30) + 1):
                v = ws.cell(row=r, column=c).value
                if v is None:
                    continue
                if comp is None and hasattr(v, 'year'):
                    comp = f'{v.year:04d}-{v.month:02d}'
                if reg is None and isinstance(v, str):
                    m = re.match(r'\s*REGI.O:?\s*(\d{1,2})', cf._norm(v))
                    if m:
                        reg = 'R' + str(int(m.group(1)))
        if reg and comp:
            break
    base = os.path.basename(nome_arquivo)
    if reg is None:
        m = re.search(r'R\.?\s*(\d{1,2})(?!\d)', base, re.IGNORECASE)
        reg = 'R' + str(int(m.group(1))) if m else None
    if comp is None:
        m = re.search(r'(JANEIRO|FEVEREIRO|MAR[CÇ]O|ABRIL|MAIO|JUNHO|JULHO|AGOSTO|SETEMBRO|OUTUBRO|NOVEMBRO|DEZEMBRO)\s*(\d{4})', base, re.IGNORECASE)
        if m:
            comp = f'{int(m.group(2)):04d}-{cf.NUM_DO_MES_PT.get(cf._norm(m.group(1))):02d}'
    return reg, comp


def ler_ficha(arquivo):
    """Lê o layout de uma ficha: grupos (colunas/rótulos) e as linhas de cada aba de trecho."""
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    abas_trecho = []
    grupos = None
    for ws in wb.worksheets:
        pos_sre = cf._achar_cabecalho(ws, 'S.R.E.', exato=True)
        if not pos_sre:
            continue
        linhas_cab = [pos_sre[0]]
        pos_ini = cf._achar_cabecalho(ws, 'INICIO', linhas=linhas_cab)
        pos_fim = cf._achar_cabecalho(ws, 'FIM', linhas=linhas_cab)
        if not (pos_ini and pos_fim):
            continue
        r_sre, c_sre = pos_sre
        inicio_dados = cf._merged_range_de(ws, r_sre, c_sre)[2] + 1
        tipo = 'pavimentada' if cf._achar_cabecalho(ws, 'CONDICAO DO PAVIMENTO', linhas=range(1, inicio_dados)) else 'nao_pavimentada'
        gs = []
        for g in cf.TEMPLATES[tipo]['grupos']:
            pos = cf._achar_cabecalho(ws, g['chave_topo'], linhas=range(1, inicio_dados))
            if not pos:
                continue
            _, cc1, r2, cc2 = cf._merged_range_de(ws, *pos)
            cols = list(range(cc1, cc2 + 1))
            rotulos = [re.sub(r'\s+', ' ', str(ws.cell(row=r2 + 1, column=c).value or '')).strip() for c in cols]
            gs.append({'id': g['id'], 'nome': g['nome'], 'rotulos': rotulos, 'cols': [get_column_letter(c) for c in cols]})
        if grupos is None:
            grupos = gs
        # nome do trecho: linha "TRECHO N – RODOVIA ..."
        tr_num, tr_nome = ws.title.strip(), ''
        for r in range(1, inicio_dados):
            v = ws.cell(row=r, column=1).value
            m = re.match(r'\s*TRECHO\s+(\d+)\s*[-–—]\s*(.+)', str(v or ''), re.IGNORECASE)
            if m:
                tr_num, tr_nome = m.group(1), m.group(2).strip()
                break
        linhas = []
        for r in range(inicio_dados, ws.max_row + 1):
            sre = re.sub(r'\s', '', str(ws.cell(row=r, column=c_sre).value or '').upper())
            ini = ws.cell(row=r, column=pos_ini[1]).value
            fim = ws.cell(row=r, column=pos_fim[1]).value
            if SRE_OK.match(sre) and isinstance(ini, (int, float)) and isinstance(fim, (int, float)):
                linhas.append((sre, round(float(ini), 3), round(float(fim), 3), r))
        if linhas:
            abas_trecho.append({'aba': ws.title, 'tr': tr_num, 'tn': tr_nome, 'linhas': linhas})
    reg, comp = regiao_e_mes([wb[a['aba']] for a in abas_trecho] or wb.worksheets, arquivo)
    return reg, comp, grupos, abas_trecho


def main():
    arquivos = sorted(glob.glob(os.path.join(FICHAS, 'fichas', 'LOTE*', '**', '*.xlsx'), recursive=True))
    melhor = {}  # (regiao, tipo) -> (competencia, arquivo, grupos, abas)
    for arq in arquivos:
        if os.path.basename(arq).startswith('~$'):
            continue
        try:
            reg, comp, grupos, abas = ler_ficha(arq)
        except Exception as erro:  # ficha com formato estranho não derruba as outras
            print('pulando', os.path.basename(arq), '-', erro)
            continue
        if not (reg and comp and abas and grupos):
            continue
        chave = (reg, tipo_do_arquivo(arq))
        if chave not in melhor or comp > melhor[chave][0]:
            melhor[chave] = (comp, arq, grupos, abas)

    pasta_modelos = os.path.join(RAIZ, 'modelos')
    os.makedirs(pasta_modelos, exist_ok=True)
    saida = []
    for (reg, tipo), (comp, arq, grupos, abas) in sorted(melhor.items(), key=lambda k: (int(k[0][0][1:]), k[0][1])):
        nome_modelo = f'ficha_{reg}_{tipo}.xlsx'
        shutil.copyfile(arq, os.path.join(pasta_modelos, nome_modelo))
        trechos = []
        n = 0
        for a in abas:
            sres = []
            for sre, ini, fim, linha_excel in a['linhas']:
                # [S.R.E., [km de cada divisa: ini, fim, fim...], [linha do Excel de cada linha]]
                if sres and sres[-1][0] == sre and abs(sres[-1][1][-1] - ini) < 1e-6:
                    sres[-1][1].append(fim)
                    sres[-1][2].append(linha_excel)
                else:
                    sres.append([sre, [ini, fim], [linha_excel]])
                n += 1
            trechos.append({'tr': a['tr'], 'tn': a['tn'], 'aba': a['aba'], 'sres': sres})
        print(f'{reg:3s} {tipo:3s} {comp}  {n:4d} linhas  {len(trechos):3d} trechos  ({os.path.basename(arq)})')
        saida.append({'r': reg, 'f': tipo, 'comp': comp, 'arq': 'modelos/' + nome_modelo, 'grupos': grupos, 'trechos': trechos})

    destino = os.path.join(RAIZ, 'js', 'modelo-ficha.js')
    with open(destino, 'w', encoding='utf-8') as f:
        f.write('// GERADO por ferramentas/gerar_modelo_ficha.py — não editar à mão.\n')
        f.write('// Linhas da Ficha de Inspeção: r=região, f=pav|np, comp=competência lida, arq=planilha-molde,\n')
        f.write('// grupos=[{id, nome, rotulos=[cabeçalho de cada coluna], cols=[letra da coluna no Excel]}],\n')
        f.write('// trechos=[{tr=nº, tn=nome, aba, sres=[[S.R.E., [km de cada divisa: 0,1,2,…,fim], [linha do Excel de cada linha]]]}]\n')
        f.write('const MODELO_FICHA = ' + json.dumps(saida, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'gravado {destino} ({os.path.getsize(destino) / 1024:.0f} KB) e {len(saida)} planilhas em modelos/')


if __name__ == '__main__':
    main()
